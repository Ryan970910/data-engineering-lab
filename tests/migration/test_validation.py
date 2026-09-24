import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from club_migration.validation import compare


class ArtifactComparison(unittest.TestCase):
    def setUp(self):
        runtime = ROOT / ".learning-runtime" / "comparison-tests"
        runtime.mkdir(parents=True, exist_ok=True)
        self.folder = tempfile.TemporaryDirectory(dir=runtime)
        self.addCleanup(self.folder.cleanup)
        self.expected = Path(self.folder.name) / "expected.json"
        self.actual = Path(self.folder.name) / "actual.json"

    def check(self, expected, actual):
        self.expected.write_text(json.dumps(expected), encoding="utf-8")
        self.actual.write_text(json.dumps(actual), encoding="utf-8")
        return compare(self.expected, self.actual)

    def test_exact_values_not_workflow_certification(self):
        result = self.check([{"id": "1", "value": None}], [{"id": "1", "value": None}])
        self.assertTrue(result["equal"])
        self.assertFalse(result["full_workflow_verified"])

    def test_same_count_different_value_fails(self):
        result = self.check([{"id": "1"}], [{"id": "2"}])
        self.assertFalse(result["equal"])

    def test_null_type_order_duplicate_and_missing_field_fail(self):
        for expected, actual in [(None, ""), (1, "1"), ([1, 2], [2, 1]), ([1, 1], [1]), ({"a": 1}, {"b": 1})]:
            with self.subTest(expected=expected):
                self.assertFalse(self.check(expected, actual)["equal"])

    def test_cannot_compare_artifact_with_itself(self):
        with self.assertRaises(ValueError):
            compare(self.expected, self.expected)

    def test_workbook_entire_schema_and_cells(self):
        from openpyxl import Workbook, load_workbook
        expected, actual = self.expected.with_suffix('.xlsx'), self.actual.with_suffix('.xlsx')
        book = Workbook()
        book.active.append(['name', 'value'])
        book.active.append(['SYNTHETIC', 2])
        book.save(expected)
        book.save(actual)
        book.close()
        self.assertTrue(compare(expected, actual)['equal'])
        book = load_workbook(actual)
        book.active['B2'] = 3
        book.save(actual)
        book.close()
        self.assertFalse(compare(expected, actual)['equal'])


if __name__ == '__main__':
    unittest.main()
