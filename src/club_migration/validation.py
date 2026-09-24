"""Strict, offline whole-artifact comparison. A match is not workflow certification."""
import argparse
import csv
import hashlib
import json
import math
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path


def _cell(value):
    if value is None:
        return ["null", None]
    if isinstance(value, float) and math.isnan(value):
        return ["float", "NaN"]
    if isinstance(value, (date, datetime, Decimal)):
        return [type(value).__name__, str(value)]
    if isinstance(value, dict):
        return ["object", [[key, _cell(item)] for key, item in value.items()]]
    if isinstance(value, (list, tuple)):
        return ["array", [_cell(item) for item in value]]
    return [type(value).__name__, value]


def _snapshot(path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=False, data_only=False)
        try:
            return {sheet.title: {
                "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
                "rows": [[[cell.data_type, cell.number_format, _cell(cell.value)] for cell in row]
                         for row in sheet.iter_rows()],
            } for sheet in book.worksheets}
        finally:
            book.close()
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return {"table": {"rows": [[_cell(v) for v in row] for row in csv.reader(stream)]}}
    if path.suffix.lower() == ".parquet" or path.is_dir():
        import pyarrow.dataset as ds

        table = ds.dataset(path, format="parquet").to_table()
        return {"table": {"schema": str(table.schema.remove_metadata()),
                          "rows": [_cell(row) for row in table.to_pylist()]}}
    if path.suffix.lower() == ".json":
        return {"json": {"rows": [_cell(json.loads(path.read_text(encoding="utf-8")))]}}
    raise ValueError("Supported artifacts: XLSX, CSV, Parquet or JSON")


def compare(expected, actual):
    expected, actual = Path(expected).resolve(), Path(actual).resolve()
    if expected == actual:
        raise ValueError("Expected and actual must be separate artifacts")
    left, right = _snapshot(expected), _snapshot(actual)
    differences = []
    if list(left) != list(right):
        differences.append({"kind": "sheet_order_or_names"})
    for name in left.keys() & right.keys():
        a, b = left[name], right[name]
        if {k: v for k, v in a.items() if k != "rows"} != {k: v for k, v in b.items() if k != "rows"}:
            differences.append({"kind": "structure", "sheet": name})
        if a["rows"] != b["rows"]:
            count = sum(x != y for x, y in zip(a["rows"], b["rows"])) + abs(len(a["rows"]) - len(b["rows"]))
            differences.append({"kind": "row_values_types_or_order", "sheet": name,
                                "expected_rows": len(a["rows"]), "actual_rows": len(b["rows"]),
                                "different_row_positions": count})
    def digest(value):
        return hashlib.sha256(json.dumps(value, ensure_ascii=False, allow_nan=False).encode("utf-8")).hexdigest()
    return {"equal": not differences, "scope": "supplied_output_artifacts_only",
            "full_workflow_verified": False, "differences": differences,
            "expected_digest": digest(left), "actual_digest": digest(right)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("expected", type=Path)
    parser.add_argument("actual", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    runtime = Path(__file__).resolve().parents[2] / ".learning-runtime"
    report = args.report.resolve()
    if not report.is_relative_to(runtime.resolve()):
        parser.error("Reports must stay in ignored .learning-runtime")
    result = compare(args.expected, args.actual)
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(report)
    raise SystemExit(0 if result["equal"] else 1)


if __name__ == "__main__":
    main()
